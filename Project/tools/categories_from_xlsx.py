"""Справочник категорий Артёма (xlsx) -> общий справочник платформы.

Запуск из папки Project:
    python tools/categories_from_xlsx.py Книга1.xlsx - книга2.xlsx \
        --dictionary=backend/contracts/dictionaries/categories.json \
        --characteristics=backend/contracts/dictionaries/characteristics.json

Второй аргумент — куда писать плоский список категорий (прежний формат редактора);
«-» — не писать. С 2026-09-25 редактор плана — модуль визарда
(frontend/src/features/planEditor) и читает только общий справочник из --dictionary,
своего categories.json у него нет.

Колонки xlsx: Название категории | К чему относится | Короткое описание | Пример.
В таблице нет id, поэтому id получается транслитерацией названия. Когда у Артёма
появятся свои id, их нужно брать из таблицы, а эту функцию убрать.

С `--dictionary=путь.json` тот же справочник дополнительно выгружается в форме,
принятой в визарде Владимирова (разделами: equipment_categories, working_zones,
point_kinds, tasks, environments) — чтобы форма и редактор читали один список.

EXTRA_EQUIPMENT_CATEGORIES (2026-09-25) — 25 категорий оборудования, которых нет
в Книга1.xlsx, но которые встретились в реальном каталоге хакатона
(`Датасет/Датасет/catalog_export_v4.csv`, колонка «Подтип», 223 позиции) — новые
семейства техники (беспилотный трамвай/метро/такси/грузовик, морские роботы и
т.п.), а не опечатки. Из 28 несовпадений 3 («Робот инвентаризатор», «Робот
уборщик», «Робот-штабелёр» без дефиса/с ё) транслитерируются в тот же id, что уже
существующие категории Книга1.xlsx — добавить их отдельными категориями нельзя
(дубль id), они не входят в этот список и просто означают ту же категорию.

Третий аргумент необязательный — книга2.xlsx (заменяет прежний артём-стас.ods,
архивирован в `../архив/`), где та же таблица Артёма разложена по укрупнённым
категориям: колонки «Обобщённая категория» | «Подкатегория» | «Исходное название
категории» | ... Из первой колонки берётся поле group: по нему редактор разбивает
оборудование в палитре. Книга2.xlsx кроме группировки содержит и новые категории,
которых нет в Книга1.xlsx (в основном БАС/морские/сельхоз/строительные задачи) —
этот скрипт их не подхватывает, т.к. список категорий по-прежнему строится из
Книга1.xlsx (src); книга2.xlsx используется только для group. Категории вида
place_zone/place_point/task/environment, которых в книга2.xlsx нет (это не
оборудование — книга2.xlsx их пока не описывает), получают группу по умолчанию
«Задачи и зоны применения» — той же, где лежат их уже сгруппированные соседи, а
не «без раздела»: в редакторе не должно быть категорий без группы.

applies_to = «характеристика» (49 строк Книга1.xlsx) — это не категория для
палитры редактора (2026-09-25: раньше они утекали в categories.json группой
«Без раздела», в редакторе не показывались — см. комментарий у kind
"characteristic" в editor2d/src/catalog/categories.ts), а атрибут оборудования:
чем гружёное/в каких условиях работает. Скрипт выносит их в отдельный список
`characteristics` и с `--characteristics=путь.json` выгружает каталог
характеристик + вручную подобранное соответствие «категория оборудования -> её
характеристики» (EQUIPMENT_CHARACTERISTICS ниже — подбор человека, не из
таблицы; механически такое соответствие не вывести).
"""

import json
import re
import sys
from datetime import datetime, timezone

import openpyxl

TRANSLIT = dict(zip(
    "абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
    ["a", "b", "v", "g", "d", "e", "e", "zh", "z", "i", "y", "k", "l", "m", "n", "o", "p",
     "r", "s", "t", "u", "f", "h", "ts", "ch", "sh", "sch", "", "y", "", "e", "yu", "ya"],
))

# Как та же категория называется в справочнике визарда (Владимиров,
# contracts/dictionaries/categories.json): id категории Артёма -> его id.
# Нужно, чтобы форма и редактор понимали друг друга: в данных может прийти
# любое из названий, а означают они одно и то же. Список сверяется при
# генерации: несуществующий id категории или повтор чужого id — ошибка.
EXTERNAL_IDS = {
    # оборудование
    "amr": ["amr"],
    "robot_shtabeler": ["stacker"],
    "robot_tyagach": ["tug"],
    "bespilotnyy_pogruzchik": ["forklift"],
    "robot_uborschik": ["cleaner"],
    "statsionarnaya_sistema_umnogo_hraneniya": ["asrs"],
    "robot_sortirovschik": ["sorter"],
    "robot_kurer": ["delivery"],
    # зоны и точки: у визарда это «рабочие зоны» формы
    "zona_priemki": ["receiving"],
    "zona_hraneniya": ["storage"],
    "zona_komplektatsii": ["picking"],
    "zona_otgruzki": ["shipping"],
    "post_upakovki": ["packing"],
    "tochka_zaryadki": ["charging"],
}

# «Зона …» из группы «точка на плане» -> тип зоны по контракту сцены (по умолчанию operation)
ZONE_TYPES = {
    "Зона хранения": "storage",
    "Зона стерильности": "restricted",
    "Зона маневрирования": "transit",
    "Зона стоянки": "transit",
    "Зона ожидания": "transit",
}

# Группа по умолчанию для place_zone/place_point/task/environment, которых нет в книга2.xlsx
# (книга2.xlsx их не описывает — это не значит, что они не относятся к этой же группе).
FALLBACK_GROUP = "Задачи и зоны применения"

# Характеристики (applies_to = «характеристика») по осям — для группировки в характеристиках.json.
WEIGHT_CLASS_IDS = {
    "tyazhelye_gruzy", "legkie_gruzy", "melkie_gruzy", "krupnogabaritnye_gruzy",
    "gruzy_do_6_kg", "gruzy_do_20_kg", "gruzy_do_40_kg", "gruzy_do_70_kg", "gruzy_do_100_kg",
    "gruzy_do_300_kg", "gruzy_do_1500_kg", "gruzy_do_2000_kg", "gruzy_do_6000_kg", "gruzy_do_35000_kg",
}
CARGO_TYPE_IDS = {
    "pallety", "telezhki", "konveyernye_moduli", "roll_keydzhi", "klt_yaschiki",
    "stellazhi", "polki", "gofrotara", "biomaterialy", "medikamenty",
}
NAVIGATION_IDS = {
    "rovnyy_pol", "razmetka_mayaki", "bez_razmetki", "gnss", "dinamicheskaya_sreda", "vysokie_stellazhi",
}
PLACE_TYPE_IDS = {
    "rabota_v_pomeschenii", "rabota_na_ulitse", "rabota_na_vysote", "rabota_v_trube", "rabota_na_zhd",
    "rabota_v_portu", "rabota_v_aeroportu", "rabota_na_sklade", "rabota_v_tsehe", "rabota_v_magazine",
    "rabota_v_bts", "rabota_v_zhk", "rabota_v_universitete", "rabota_v_bolnitse", "rabota_v_laboratorii",
    "rabota_na_obekte_transporta_i_logistiki", "rabota_na_obekte_bezopasnosti",
    "rabota_na_obekte_torgovli_i_uslug", "rabota_na_obekte_promyshlennosti",
}
CHARACTERISTIC_AXES = {
    cid: axis
    for ids, axis in [
        (WEIGHT_CLASS_IDS, "weight_class"),
        (CARGO_TYPE_IDS, "cargo_type"),
        (NAVIGATION_IDS, "navigation"),
        (PLACE_TYPE_IDS, "place_type"),
    ]
    for cid in ids
}

# id категории оборудования -> её характеристики (id из CHARACTERISTIC_AXES). Подобрано вручную
# по описанию/назначению категории (2026-09-25) — не выводится из таблиц Артёма механически;
# при появлении новой техники в Книга1.xlsx новой категории нужно завести запись здесь.
EQUIPMENT_CHARACTERISTICS: dict[str, list[str]] = {
    "mobilnyy_robot": ["legkie_gruzy", "pallety", "telezhki", "rovnyy_pol", "rabota_na_sklade"],
    "amr": ["legkie_gruzy", "telezhki", "roll_keydzhi", "bez_razmetki", "rovnyy_pol",
            "rabota_na_sklade", "rabota_v_tsehe", "rabota_v_magazine"],
    "fmr_avtonomnyy_pogruzchik": ["tyazhelye_gruzy", "gruzy_do_2000_kg", "pallety", "stellazhi",
                                   "razmetka_mayaki", "rovnyy_pol", "rabota_na_sklade", "rabota_v_tsehe"],
    "robot_shtabeler": ["tyazhelye_gruzy", "pallety", "stellazhi", "vysokie_stellazhi",
                         "razmetka_mayaki", "rovnyy_pol", "rabota_na_sklade"],
    "robot_tyagach": ["tyazhelye_gruzy", "gruzy_do_2000_kg", "telezhki", "roll_keydzhi",
                       "razmetka_mayaki", "rovnyy_pol", "rabota_na_sklade", "rabota_v_tsehe"],
    "robot_uborschik": ["rovnyy_pol", "dinamicheskaya_sreda", "rabota_v_pomeschenii", "rabota_na_ulitse",
                         "rabota_na_sklade", "rabota_v_magazine", "rabota_v_bts"],
    "bespilotnyy_pogruzchik": ["tyazhelye_gruzy", "gruzy_do_2000_kg", "pallety", "razmetka_mayaki",
                                "rovnyy_pol", "rabota_na_sklade", "rabota_v_portu"],
    "statsionarnaya_sistema_umnogo_hraneniya": ["melkie_gruzy", "klt_yaschiki", "stellazhi", "rabota_na_sklade"],
    "robot_manipulyator": ["tyazhelye_gruzy", "dinamicheskaya_sreda", "rabota_v_tsehe",
                            "rabota_na_obekte_promyshlennosti"],
    "robot_ukladchik": ["tyazhelye_gruzy", "gofrotara", "stellazhi", "rabota_v_tsehe"],
    "delta_robot": ["melkie_gruzy", "konveyernye_moduli", "rabota_v_tsehe", "rabota_na_obekte_promyshlennosti"],
    "kollaborativnyy_robot_kobot": ["legkie_gruzy", "dinamicheskaya_sreda", "rabota_v_tsehe",
                                     "rabota_na_obekte_promyshlennosti"],
    "robot_sortirovschik": ["legkie_gruzy", "melkie_gruzy", "roll_keydzhi", "konveyernye_moduli",
                             "razmetka_mayaki", "rovnyy_pol", "rabota_na_sklade",
                             "rabota_na_obekte_transporta_i_logistiki"],
    "robot_inventarizator": ["bez_razmetki", "vysokie_stellazhi", "rabota_na_sklade", "rabota_v_magazine"],
    "robot_rover": ["legkie_gruzy", "gnss", "rabota_na_ulitse", "rabota_v_zhk", "rabota_v_universitete"],
    "robot_kurer": ["legkie_gruzy", "gnss", "rabota_na_ulitse", "rabota_v_zhk"],
    "antropomorfnyy_robot": ["legkie_gruzy", "dinamicheskaya_sreda", "rabota_v_bts", "rabota_v_magazine",
                              "rabota_v_universitete"],
    "robo_kafe": ["melkie_gruzy", "rabota_v_bts", "rabota_v_magazine", "rabota_na_obekte_torgovli_i_uslug"],
    "mobilnyy_manipulyator": ["legkie_gruzy", "tyazhelye_gruzy", "dinamicheskaya_sreda", "rovnyy_pol",
                               "rabota_v_tsehe", "rabota_na_sklade"],
    "robotizirovannaya_ustanovka": ["tyazhelye_gruzy", "rabota_na_obekte_promyshlennosti"],
    "robotizirovannaya_pristanochnaya_yacheyka": ["tyazhelye_gruzy", "rabota_v_tsehe"],
    "robotizirovannyy_svarochnyy_kompleks": ["tyazhelye_gruzy", "rabota_v_tsehe",
                                              "rabota_na_obekte_promyshlennosti"],
    "robotizirovannye_pogruzochnye_krany": ["tyazhelye_gruzy", "gruzy_do_35000_kg", "rabota_v_portu",
                                             "rabota_na_obekte_promyshlennosti"],
    "vnutritrubnaya_robotizirovannaya_sistema": ["bez_razmetki", "rabota_v_trube"],
    "robot_laborant": ["melkie_gruzy", "biomaterialy", "medikamenty", "rabota_v_laboratorii", "rabota_v_bolnitse"],
    "robot_ekolog_obhodchik": ["legkie_gruzy", "gnss", "bez_razmetki", "rabota_na_ulitse",
                                "rabota_na_obekte_bezopasnosti"],
    "ohrannyy_robot": ["legkie_gruzy", "gnss", "bez_razmetki", "dinamicheskaya_sreda",
                        "rabota_na_obekte_bezopasnosti", "rabota_na_ulitse", "rabota_na_sklade"],
    "robotizirovannyy_lafetnyy_stvol": ["rabota_na_obekte_bezopasnosti", "rabota_na_obekte_promyshlennosti"],
    "robotizirovannaya_telezhka": ["legkie_gruzy", "telezhki", "razmetka_mayaki", "rovnyy_pol",
                                    "rabota_v_tsehe", "rabota_na_sklade"],
    "universalnaya_avtonomnaya_robotizirovannaya_platforma": ["legkie_gruzy", "tyazhelye_gruzy", "bez_razmetki",
                                                                "gnss", "rabota_v_pomeschenii", "rabota_na_ulitse"],
    "bas": ["legkie_gruzy", "gruzy_do_20_kg", "gnss", "rabota_na_ulitse"],
    "bas_multirotornogo_tipa": ["gruzy_do_6_kg", "gruzy_do_20_kg", "gnss", "rabota_na_ulitse"],
    "bas_samoletnogo_tipa": ["gruzy_do_6_kg", "gnss", "rabota_na_ulitse"],
    "vtol": ["gruzy_do_20_kg", "gruzy_do_40_kg", "gnss", "rabota_na_ulitse"],
    "bas_dlya_aerofotosemki": ["gruzy_do_6_kg", "gnss", "rabota_na_ulitse", "rabota_na_obekte_promyshlennosti"],
    "bas_dlya_monitoringa": ["gruzy_do_6_kg", "gnss", "rabota_na_ulitse", "rabota_na_obekte_bezopasnosti"],
    "bas_dlya_dostavki_gruzov": ["gruzy_do_20_kg", "gruzy_do_40_kg", "gnss", "rabota_na_ulitse"],
    "bas_dlya_moyki_fasadov": ["gruzy_do_20_kg", "gnss", "rabota_na_vysote", "rabota_na_ulitse"],
    "bas_dlya_poiska_propavshih": ["gruzy_do_6_kg", "gnss", "rabota_na_ulitse", "rabota_na_obekte_bezopasnosti"],
    # --- EXTRA_EQUIPMENT_CATEGORIES (catalog_export_v4.csv, см. ниже) ---
    "fmr": ["tyazhelye_gruzy", "pallety", "rovnyy_pol", "razmetka_mayaki", "rabota_na_sklade"],
    "bespilotnoe_metro": ["rabota_na_zhd", "gnss"],
    "bespilotnoe_taksi": ["gnss", "rabota_na_ulitse"],
    "bespilotnyy_gruzovik": ["tyazhelye_gruzy", "gruzy_do_2000_kg", "gnss", "rabota_na_ulitse",
                              "rabota_na_obekte_transporta_i_logistiki"],
    "bespilotnyy_tramvay": ["rabota_na_zhd", "gnss"],
    "bespilotnyy_tyagach": ["tyazhelye_gruzy", "telezhki", "gnss", "rabota_na_ulitse"],
    "intellektualnyy_avtobus_vendingovyy_apparat": ["gnss", "rabota_na_ulitse", "rabota_na_obekte_torgovli_i_uslug"],
    "vnutritrubnye_robotizirovannye_sistemy": ["bez_razmetki", "rabota_v_trube"],
    "mobilnyy_robot_komplektovschik": ["legkie_gruzy", "roll_keydzhi", "rovnyy_pol", "rabota_na_sklade"],
    "robot_rastsepschik": ["rabota_na_zhd"],
    "robot": ["rabota_na_zhd"],
    "rover": ["rabota_na_zhd"],
    "kollaborativnyy_robot_manipulyator": ["legkie_gruzy", "dinamicheskaya_sreda", "rabota_v_tsehe"],
    "manipulyator": ["tyazhelye_gruzy", "rabota_v_tsehe"],
    "kran_shtabeler": ["tyazhelye_gruzy", "pallety", "stellazhi", "rabota_na_sklade"],
    "robotizirovannyy_3d_printer": ["rabota_na_obekte_promyshlennosti"],
    "umnaya_sistema_hraneniya": ["melkie_gruzy", "klt_yaschiki", "stellazhi", "rabota_na_sklade"],
    "shattl": ["tyazhelye_gruzy", "pallety", "stellazhi", "rabota_na_sklade"],
    "bezekipazhnyy_kater": ["rabota_na_ulitse"],
    "bespilotnyy_katamaran": ["rabota_na_ulitse"],
    "mnogofunktsionalnyy_bezekipazhnyy_kater": ["rabota_na_ulitse"],
    "modulnaya_navodnaya_mnogofunktsionalnaya_platforma": ["rabota_na_ulitse"],
    "multirotor": ["gruzy_do_6_kg", "gruzy_do_20_kg", "gnss", "rabota_na_ulitse"],
    "privyaznoy_robot_na_baze_oktokoptera": ["gnss", "rabota_na_vysote", "rabota_na_ulitse"],
    "samolet": ["gruzy_do_6_kg", "gnss", "rabota_na_ulitse"],
}

# 25 категорий оборудования из реального каталога (catalog_export_v4.csv, колонка «Подтип»),
# которых нет в Книга1.xlsx — см. докстринг модуля. (name, group, description, examples).
EXTRA_EQUIPMENT_CATEGORIES: list[tuple[str, str, str, list[str]]] = [
    ("FMR", "Мобильные роботы",
     "Мобильный вилочный робот (класс FMR) — автономное напольное ТС для паллетных грузов",
     ["AK-2000-2"]),
    ("Беспилотное метро", "Мобильные роботы", "Автономный рельсовый транспорт метрополитена", ["Беспилотный поезд «2024»"]),
    ("Беспилотное такси", "Мобильные роботы", "Автономный легковой автомобиль для перевозки пассажиров", ["Беспилотное такси Яндекс"]),
    ("Беспилотный грузовик", "Мобильные роботы", "Автономный электрический грузовик вне объекта", ["EVOCARGO N1"]),
    ("Беспилотный трамвай", "Мобильные роботы", "Автономный рельсовый городской транспорт", ["Беспилотный «Львёнок-Москва»"]),
    ("Беспилотный тягач", "Мобильные роботы", "Автономный тягач для буксировки тележек/составов вне здания", ["Робот-тягач RoboCV"]),
    ("Интеллектуальный автобус — вендинговый аппарат", "Мобильные роботы",
     "Автономный мобильный вендинговый автомат на колёсной платформе", ["IQ-BUS V"]),
    ("Внутритрубные роботизированные системы", "Мобильные роботы",
     "Роботы для диагностики/очистки труднопроходимых (unpiggable) трубопроводов", ["Тьюбот"]),
    ("Мобильный робот-комплектовщик", "Мобильные роботы", "Мобильный робот для комплектации заказов на складе", ["Робот-комплектовщик"]),
    ("Робот-расцепщик", "Мобильные роботы", "Робототехнический комплекс расцепки вагонов на сортировочных горках", ["Робот-расцепщик"]),
    ("Робот", "Мобильные роботы", "Без уточнения подкласса в каталоге; пример — комплекс отпуска тормозов вагонов", ["Робот-растормаживатель"]),
    ("Ровер", "Мобильные роботы", "Колёсный ровер; в каталоге — для работы с грузовыми вагонами", ["Ровер для работы с грузовыми вагонами"]),
    ("Коллаборативный робот-манипулятор", "Стационарные роботы и комплексы",
     "Манипулятор, работающий совместно с человеком без ограждения", ["RC5-16"]),
    ("Манипулятор", "Стационарные роботы и комплексы", "Промышленный манипулятор без уточнения подкласса", ["Роборука"]),
    ("Кран-штабелёр", "Стационарные роботы и комплексы", "Автоматический кран для стеллажного хранения паллет", ["AS-RS P"]),
    ("Роботизированный 3D принтер", "Стационарные роботы и комплексы",
     "Промышленный робот для 3D-печати строительных элементов", ["АМТ S300"]),
    ("Умная система хранения", "Стационарные роботы и комплексы",
     "Роботизированная система кубического хранения (AutoStore-подобная)", ["SmartCube"]),
    ("Шаттл", "Стационарные роботы и комплексы", "Шаттл для стеллажной системы высокой плотности хранения", ["Pallet Shuttle"]),
    ("Безэкипажный катер", "Морские и подводные роботы", "Автономное надводное судно", ["Сарган"]),
    ("Беспилотный катамаран", "Морские и подводные роботы", "Автономное надводное судно катамаранного типа", ["ОМДЖЕТ-ГМ01"]),
    ("Многофункциональный бэзэкипажный катер", "Морские и подводные роботы",
     "Многофункциональное автономное надводное судно", ["БЭК-Р"]),
    ("Модульная наводная многофункциональная платформа", "Морские и подводные роботы",
     "Модульная автономная надводная платформа", ["Барабулька"]),
    ("Мультиротор", "Воздушные роботы (БАС)", "БАС мультироторного типа без уточнения назначения", ["А-50"]),
    ("Привязной робот на базе октокоптера", "Воздушные роботы (БАС)",
     "Привязной БАС для мойки фасадов высотных зданий", ["Astramis AeroUnit"]),
    ("Самолет", "Воздушные роботы (БАС)", "БАС самолётного типа без уточнения назначения", ["Геоскан 201"]),
]


def make_id(name: str) -> str:
    s = "".join(TRANSLIT.get(ch, ch) for ch in name.lower())
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def classify(name: str, applies_to: str, description: str) -> str:
    """Во что категория превращается в редакторе."""
    a = applies_to.lower()
    if a == "оборудование":
        return "software" if name.startswith("ПО") else "equipment"
    if a == "характеристика":
        return "characteristic"
    if a == "точка на плане":
        return "place_zone" if name.startswith("Зона") else "place_point"
    # «точка на плане / и то и другое»: условия среды описаны как «Работа …», остальное — задачи
    return "environment" if description.startswith("Работа") else "task"


def read_groups(path: str) -> dict[str, str]:
    """Название категории -> обобщённая категория (книга2.xlsx). Колонки:
    Обобщённая категория | Подкатегория | Исходное название категории | ..."""
    ws = openpyxl.load_workbook(path, data_only=True).active
    groups: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[2]:
            continue
        groups.setdefault(str(row[2]).strip(), str(row[0]).strip())
    if not groups:
        sys.exit(f"в {path} не нашлось ни одной категории с группой — проверь формат файла")
    return groups


#: Разделы справочника визарда: как он называет то, что у Артёма отличается видом (kind)
DICTIONARY_SECTIONS = [
    ("equipment_categories", "equipment"),
    ("working_zones", "place_zone"),
    ("point_kinds", "place_point"),
    ("tasks", "task"),
    ("environments", "environment"),
]


def write_dictionary(categories: list[dict], dst: str) -> None:
    """Тот же справочник в форме, принятой в визарде (contracts/dictionaries/categories.json):
    разделами, с главным id и списком других названий той же категории.

    Главный id — английский, если он уже заведён у Владимирова (тогда наш транслит уходит
    в aliases); у остальных категорий главным остаётся транслит из таблицы Артёма, потому
    что английского названия для них никто не утверждал — придумывать его скрипту нельзя.
    """
    out: dict = {
        "$comment": (
            "Собран скриптом tools/categories_from_xlsx.py из таблицы Артёма — руками не править. "
            "id — главное название категории, aliases — как её же называют в другом справочнике; "
            "и то и другое означает одну категорию. group — раздел таблицы Артёма."
        ),
        "version": 3,
        "source": "Книга1.xlsx + книга2.xlsx",
    }
    for section, kind in DICTIONARY_SECTIONS:
        items = []
        for c in categories:
            if c["kind"] != kind:
                continue
            external = c["aliases"]
            main_id = external[0] if external else c["id"]
            other = [a for a in ([c["id"]] + external[1:] if external else []) if a != main_id]
            item = {"id": main_id, "label": c["name"], "aliases": other, "group": c["group"]}
            if kind == "place_zone":
                item["zone_type"] = c["zone_type"]
            if kind == "place_point":
                item["point_kind"] = c["point_kind"]
            items.append(item)
        out[section] = items
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"справочник для визарда -> {dst}: " + ", ".join(f"{s} {len(out[s])}" for s, _ in DICTIONARY_SECTIONS))


def write_characteristics(characteristics: list[dict], dst: str) -> None:
    """Каталог характеристик оборудования + EQUIPMENT_CHARACTERISTICS (характеристики.json,
    для catalog/matching — не для редактора: editor2d характеристики не показывает)."""
    known = {c["id"] for c in characteristics}
    unknown_refs = {cid for ids in EQUIPMENT_CHARACTERISTICS.values() for cid in ids} - known
    if unknown_refs:
        sys.exit(f"в EQUIPMENT_CHARACTERISTICS есть характеристики, которых нет в таблице: {sorted(unknown_refs)}")
    out = {
        "$comment": (
            "Собран скриптом tools/categories_from_xlsx.py из таблицы Артёма (applies_to = "
            "«характеристика») — руками не править. characteristics — каталог, axis — ось "
            "группировки (weight_class/cargo_type/navigation/place_type). "
            "equipment_characteristics — id категории оборудования (см. categories.json) -> её "
            "характеристики; подобрано вручную (EQUIPMENT_CHARACTERISTICS в скрипте), не из таблицы."
        ),
        "version": 1,
        "source": "Книга1.xlsx",
        "characteristics": characteristics,
        "equipment_characteristics": EQUIPMENT_CHARACTERISTICS,
    }
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"{len(characteristics)} характеристик, {len(EQUIPMENT_CHARACTERISTICS)} категорий с назначением -> {dst}")


def main(src: str, dst: str, groups_src: str | None = None, dictionary_dst: str | None = None,
         characteristics_dst: str | None = None) -> None:
    groups = read_groups(groups_src) if groups_src else {}
    ws = openpyxl.load_workbook(src, data_only=True).active
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]

    categories, characteristics, seen = [], [], set()
    for name, applies_to, description, examples in rows:
        name, applies_to = str(name).strip(), str(applies_to or "").strip()
        description = str(description or "").strip()
        cid = make_id(name)
        if cid in seen:
            sys.exit(f"повторяющийся id {cid!r} у категории {name!r}")
        seen.add(cid)
        examples_list = [e.strip() for e in str(examples or "").split(",") if e.strip()]

        kind = classify(name, applies_to, description)
        if kind == "characteristic":
            # атрибут оборудования, не категория для палитры редактора — отдельный список
            characteristics.append({
                "id": cid,
                "name": name,
                "axis": CHARACTERISTIC_AXES.get(cid),
                "description": description,
                "examples": examples_list,
            })
            continue

        group = groups.get(name)
        if group is None and kind in ("task", "place_point", "place_zone", "environment"):
            group = FALLBACK_GROUP  # книга2.xlsx пока не описывает эту категорию — не «без раздела»
        item = {
            "id": cid,
            "name": name,
            "kind": kind,
            "group": group,
            "aliases": EXTERNAL_IDS.get(cid, []),  # те же категории в справочнике визарда
            "applies_to": applies_to,
            "description": description,
            "examples": examples_list,
        }
        if kind == "place_zone":
            item["zone_type"] = ZONE_TYPES.get(name, "operation")
        if kind == "place_point":
            item["point_kind"] = "charging" if name == "Точка зарядки" else "operation"
        categories.append(item)

    for name, group, description, examples_list in EXTRA_EQUIPMENT_CATEGORIES:
        cid = make_id(name)
        if cid in seen:
            sys.exit(f"EXTRA_EQUIPMENT_CATEGORIES: id {cid!r} ({name!r}) уже есть в Книга1.xlsx — не дублируй, это одна категория")
        seen.add(cid)
        categories.append({
            "id": cid,
            "name": name,
            "kind": "equipment",
            "group": group,
            "aliases": EXTERNAL_IDS.get(cid, []),
            "applies_to": "оборудование",
            "description": description,
            "examples": examples_list,
            "source": "catalog_export_v4.csv",  # не из Книга1.xlsx — для прозрачности происхождения
        })

    # сверяем таблицу соответствий: опечатка в ней тихо разорвала бы связь формы и редактора
    unknown = [cid for cid in EXTERNAL_IDS if cid not in seen]
    if unknown:
        sys.exit(f"в EXTERNAL_IDS есть id, которых нет в таблице: {unknown}")
    used: dict[str, str] = {}
    for cid, external in EXTERNAL_IDS.items():
        for e in external:
            if e in used:
                sys.exit(f"чужой id {e!r} указан сразу у {used[e]!r} и {cid!r}")
            used[e] = cid

    equipment_ids = {c["id"] for c in categories if c["kind"] == "equipment"}
    stale = set(EQUIPMENT_CHARACTERISTICS) - equipment_ids
    if stale:
        sys.exit(f"в EQUIPMENT_CHARACTERISTICS есть id категорий, которых нет среди оборудования: {sorted(stale)}")
    unassigned = equipment_ids - set(EQUIPMENT_CHARACTERISTICS)
    if unassigned:
        sys.exit(f"в Книга1.xlsx есть оборудование без записи в EQUIPMENT_CHARACTERISTICS: {sorted(unassigned)}")

    out = {
        "source": src.replace("\\", "/").split("/")[-1],
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "categories": categories,
    }
    if dst != "-":
        with open(dst, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False, indent=2)
            f.write("\n")

    counts: dict[str, int] = {}
    for c in categories:
        counts[c["kind"]] = counts.get(c["kind"], 0) + 1
    print(f"{len(categories)} категорий{'' if dst == '-' else f' -> {dst}'}: {counts}")
    if dictionary_dst:
        write_dictionary(categories, dictionary_dst)
    if characteristics_dst:
        write_characteristics(characteristics, characteristics_dst)
    if groups:
        no_group = [c["name"] for c in categories if not c["group"]]
        print(f"разделов: {len(set(groups.values()))}; без раздела: {no_group or 'нет'}")


if __name__ == "__main__":
    flags = {"--dictionary": None, "--characteristics": None}
    args = []
    for a in sys.argv[1:]:
        matched = next((f for f in flags if a.startswith(f + "=")), None)
        if matched:
            flags[matched] = a.split("=", 1)[1]
        else:
            args.append(a)
    if len(args) not in (2, 3):
        sys.exit(__doc__)
    main(args[0], args[1], args[2] if len(args) == 3 else None, flags["--dictionary"], flags["--characteristics"])
